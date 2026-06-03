from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from datetime import datetime
import customtkinter as ctk
from tkinter import ttk, messagebox, filedialog, StringVar, END, BOTH, RIGHT, LEFT, Y, X, BOTTOM
import sqlite3
import pandas as pd
import os
import hashlib
import time
import traceback
import threading
import queue
import multiprocessing
import gc  # Manajemen memori
import urllib.request
import json
import webbrowser

from tkcalendar import Calendar

os.environ["PYTHONOPTIMIZE"] = "1"

APP_NAME = "RESI SCANNER PAPAYA"
VERSION = "1.0.0"  # Versi aplikasi saat ini
DB_FOLDER = "databases"
BACKUP_FOLDER = "backup"
USER_DB = "users.db"

# Mapping 7 Pilihan Database
DB_MAPPING = {
    "ID express": "id_express.db",
    "JNT (cargo)": "jnt_cargo.db",
    "JNE": "jne.db",
    "Anteraja": "anteraja.db",
    "Shopee": "shopee.db",
    "Paxel": "paxel.db",
    "Other": "other.db"
}

os.makedirs(DB_FOLDER, exist_ok=True)
os.makedirs(BACKUP_FOLDER, exist_ok=True)

ctk.set_appearance_mode("light")
ctk.set_default_color_theme("blue")

def hash_password(password):
    return hashlib.sha256(password.encode()).hexdigest()

class App:
    def __init__(self, root):
        self.root = root
        self.root.title(APP_NAME)

        try:
            self.root.state("zoomed")
        except Exception:
            self.root.attributes("-zoomed", True)

        self.root.protocol("WM_DELETE_WINDOW", self.exit_app)

        self.conn = None
        self.cursor = None
        self.user_conn = None
        self.user_cursor = None
        self.current_db = ""

        self.current_user = ""
        self.role = ""
        self.last_scan_time = 0
        self.scan_queue = queue.Queue()
        self.processing_scan = False

        self.popup_open = False
        self.is_running = True
        self.scan_job = None
        self.scan_entry = None  # Inisialisasi awal agar hasattr aman
        self.current_date_filter = ""
        
        self.total_count = 0
        self.normal_count = 0
        self.double_count = 0
        
        self.session_authorized_dates = set()

        self.connect_user_db()
        self.login_ui()

    def clear(self):
        for widget in self.root.winfo_children():
            widget.destroy()
        gc.collect()

    def popup_config(self, win):
        self.popup_open = True
        win.lift()
        win.after(200, lambda: win.attributes("-topmost", True))
        win.after(500, lambda: win.attributes("-topmost", False))

        def on_close():
            self.popup_open = False
            try:
                win.destroy()
            except Exception:
                pass
            gc.collect()
        win.protocol("WM_DELETE_WINDOW", on_close)

    def connect_user_db(self):
        self.user_conn = sqlite3.connect(USER_DB)
        self.user_cursor = self.user_conn.cursor()
        self.user_cursor.execute("""
            CREATE TABLE IF NOT EXISTS users(
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT UNIQUE,
                password TEXT,
                role TEXT
            )
        """)
        # Fitur Baru: Tabel untuk menyimpan catatan/reminder antar-user
        self.user_cursor.execute("""
            CREATE TABLE IF NOT EXISTS reminders(
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT,
                pesan TEXT,
                waktu TEXT
            )
        """)
        self.user_conn.commit()
        
        accounts = [
            ("admin", "admin123", "admin"),
            ("user", "user123", "user"),
            ("operator", "operator123", "operator")
        ]
        for username, password, role in accounts:
            self.user_cursor.execute("SELECT id FROM users WHERE username=?", (username,))
            if not self.user_cursor.fetchone():
                self.user_cursor.execute(
                    "INSERT INTO users(username,password,role) VALUES(?,?,?)",
                    (username, hash_password(password), role)
                )
        self.user_conn.commit()

    def connect_db(self):
        try:
            if not self.current_db:
                raise Exception("Database belum dipilih")

            if self.conn:
                self.conn.close()

            self.conn = sqlite3.connect(self.current_db, timeout=30, check_same_thread=False)
            self.cursor = self.conn.cursor()
            
            self.cursor.execute("PRAGMA journal_mode=WAL")
            self.cursor.execute("PRAGMA synchronous=NORMAL")
            self.cursor.execute("PRAGMA cache_size=-2000")
            self.cursor.execute("PRAGMA temp_store=MEMORY")
            
            self.cursor.execute("""
                CREATE TABLE IF NOT EXISTS scans(
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    barcode TEXT,
                    status TEXT,
                    username TEXT,
                    waktu TEXT,
                    unique_id TEXT
                )
            """)
            self.cursor.execute("CREATE INDEX IF NOT EXISTS idx_barcode ON scans(barcode)")
            self.cursor.execute("CREATE TABLE IF NOT EXISTS locked_dates(date TEXT PRIMARY KEY)")
            
            self.cursor.execute("""
                CREATE TABLE IF NOT EXISTS logs(
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    username TEXT,
                    activity TEXT,
                    waktu TEXT
                )
            """)
            self.conn.commit()
            return True
        except Exception as e:
            messagebox.showerror("DATABASE ERROR", str(e))
            return False

    def add_log(self, activity):
        try:
            if not self.cursor: return
            self.cursor.execute(
                "INSERT INTO logs(username,activity,waktu) VALUES(?,?,?)",
                (self.current_user, activity, datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
            )
            self.conn.commit()
        except Exception as e:
            print(f"Log Error: {e}")

    def login_ui(self):
        self.clear()
        frame = ctk.CTkFrame(self.root, fg_color="transparent")
        frame.pack(expand=True)

        ctk.CTkLabel(frame, text="LOGIN SYSTEM", font=("Segoe UI", 24, "bold")).pack(pady=20)

        self.username = ctk.CTkEntry(frame, font=("Segoe UI", 14), width=250, placeholder_text="Username")
        self.username.pack(pady=10)

        self.password = ctk.CTkEntry(frame, font=("Segoe UI", 14), width=250, show="*", placeholder_text="Password")
        self.password.pack(pady=10)
        self.password.bind("<Return>", lambda e: self.validate_login())

        ctk.CTkButton(frame, text="LOGIN", font=("Segoe UI", 12, "bold"), width=250, height=40, command=self.validate_login).pack(pady=20)
        ctk.CTkButton(frame, text="EXIT APP", font=("Segoe UI", 12, "bold"), width=250, height=40, fg_color="#D32F2F", hover_color="#B71C1C", command=self.exit_app).pack(pady=5)

        about_btn = ctk.CTkButton(self.root, text="ℹ ABOUT", width=80, height=30, fg_color="#37474F", hover_color="#455A64", text_color="white", font=("Segoe UI", 10, "bold"), corner_radius=8, command=self.show_about_popup)
        about_btn.pack(side=BOTTOM, anchor="e", padx=20, pady=20)

    def validate_login(self):
        try:
            user = self.username.get().strip().lower()
            pw = hash_password(self.password.get())

            self.user_cursor.execute("SELECT * FROM users WHERE username=? AND password=?", (user, pw))
            data = self.user_cursor.fetchone()

            if not data:
                messagebox.showerror("LOGIN", "Username atau password salah")
                return

            self.current_user = data[1]
            self.role = data[3]
            self.database_selector_gui()
        except Exception as e:
            messagebox.showerror("LOGIN ERROR", str(e))

    def database_selector_gui(self):
        self.reset_session()
        win = ctk.CTkToplevel(self.root)
        win.title("DATABASE COURIER SELECTOR")
        win.geometry("450x350")
        self.popup_config(win)

        ctk.CTkLabel(win, text="PILIH DATABASE COURIER", font=("Segoe UI", 20, "bold")).pack(pady=20)

        db_options = list(DB_MAPPING.keys())
        self.db_choice = StringVar(value=db_options[0])
        
        dropdown = ctk.CTkOptionMenu(win, variable=self.db_choice, values=db_options, width=280, height=40)
        dropdown.pack(pady=15)

        def load_selected_db():
            selected = self.db_choice.get()
            filename = DB_MAPPING.get(selected, "other.db")
            self.current_db = os.path.join(DB_FOLDER, filename)
            
            if self.connect_db():
                self.process_login(win)

        ctk.CTkButton(win, text="LOAD & MASUK DATABASE", font=("Segoe UI", 12, "bold"), width=280, height=45, fg_color="#388E3C", hover_color="#2E7D32", command=load_selected_db).pack(pady=15)
        ctk.CTkButton(win, text="GANTI USER", font=("Segoe UI", 12, "bold"), width=280, height=40, fg_color="#E65100", hover_color="#F57C00", command=lambda: [win.destroy(), self.logout()]).pack(pady=5)

    def process_login(self, win):
        try:
            self.add_log("LOGIN")
            win.destroy()
            self.current_date_filter = datetime.now().strftime("%Y-%m-%d")
            self.main_ui()
        except Exception as e:
            messagebox.showerror("LOGIN ERROR", str(e))

    def main_ui(self):
        self.is_running = True
        self.clear()

        top = ctk.CTkFrame(self.root, fg_color="transparent")
        top.pack(fill=X, padx=10, pady=10)

        active_db_name = os.path.basename(self.current_db).replace('.db', '').upper()
        ctk.CTkLabel(top, text=f"USER : {self.current_user} ({self.role.upper()})  |  DATABASE AKTIF: {active_db_name}", font=("Segoe UI", 14, "bold")).pack(side=LEFT)

        btn_kwargs = {"height": 35, "width": 100, "font": ("Segoe UI", 11, "bold")}
        ctk.CTkButton(top, text="LOGOUT", fg_color="#D32F2F", hover_color="#B71C1C", command=self.logout, **btn_kwargs).pack(side=RIGHT, padx=5)
        ctk.CTkButton(top, text="CLOSE DB", command=self.close_database, **btn_kwargs).pack(side=RIGHT, padx=5)

        if self.role == "admin":
            ctk.CTkButton(top, text="BACKUP DB", fg_color="#00695C", hover_color="#004D40", command=self.backup_database, **btn_kwargs).pack(side=RIGHT, padx=5)
            
        if self.role in ["admin", "user"]:
            ctk.CTkButton(top, text="PDF", command=self.export_pdf, **btn_kwargs).pack(side=RIGHT, padx=5)
            ctk.CTkButton(top, text="EXCEL", command=self.export_excel, **btn_kwargs).pack(side=RIGHT, padx=5)
            ctk.CTkButton(top, text="SEARCH", command=self.search_data, **btn_kwargs).pack(side=RIGHT, padx=5)

        if self.role == "admin":
            admin_frame = ctk.CTkFrame(self.root, fg_color="transparent")
            admin_frame.pack(fill=X, padx=10, pady=5)
            ctk.CTkButton(admin_frame, text="ADD USER", command=self.add_user, **btn_kwargs).pack(side=RIGHT, padx=5)
            ctk.CTkButton(admin_frame, text="GANTI PW", fg_color="#8E24AA", hover_color="#6A1B9A", command=self.change_password, **btn_kwargs).pack(side=RIGHT, padx=5)
            ctk.CTkButton(admin_frame, text="VIEW LOG", command=self.view_logs, **btn_kwargs).pack(side=RIGHT, padx=5)
            ctk.CTkButton(admin_frame, text="EDIT RESI", command=self.edit_barcode, **btn_kwargs).pack(side=RIGHT, padx=5)
            ctk.CTkButton(admin_frame, text="DELETE", fg_color="#F57C00", hover_color="#EF6C00", command=self.delete_data, **btn_kwargs).pack(side=RIGHT, padx=5)
            ctk.CTkButton(admin_frame, text="IMPORT", command=self.import_excel, **btn_kwargs).pack(side=RIGHT, padx=5)

        if self.role in ["admin", "operator"]:
            mid = ctk.CTkFrame(self.root, fg_color="transparent")
            mid.pack(fill=X, pady=10)
            ctk.CTkLabel(mid, text="MASUKAN BARCODE", font=("Segoe UI", 16, "bold")).pack()
            
            self.scan_entry = ctk.CTkEntry(mid, font=("Segoe UI", 24), justify="center", height=50)
            self.scan_entry.pack(fill=X, padx=20, pady=10)
            self.scan_entry.bind("<Return>", lambda e: self.scan_barcode())

        info_frame = ctk.CTkFrame(self.root, fg_color="transparent")
        info_frame.pack(fill=X, padx=20, pady=5)
        
        ctk.CTkLabel(info_frame, text="Filter Tanggal:", font=("Segoe UI", 12, "bold")).pack(side=LEFT, padx=5)
        
        self.date_lbl = ctk.CTkLabel(info_frame, text=self.current_date_filter if self.current_date_filter else "Semua Data", font=("Segoe UI", 12), text_color="blue")
        self.date_lbl.pack(side=LEFT, padx=10)
        
        ctk.CTkButton(info_frame, text="📅 PILIH KALENDER", width=130, height=30, font=("Segoe UI", 11), command=self.open_calendar_filter).pack(side=LEFT, padx=5)
        ctk.CTkButton(info_frame, text="RESET FILTER", width=100, height=30, fg_color="gray", font=("Segoe UI", 11), command=self.reset_date_filter).pack(side=LEFT, padx=5)
        
        # Fitur Baru: Tombol Akses Menu Catatan/Reminder Bersama untuk semua jenis User
        ctk.CTkButton(info_frame, text="📝 CATATAN / REMINDER", width=150, height=30, fg_color="#00838F", hover_color="#006064", font=("Segoe UI", 11, "bold"), command=self.open_reminders).pack(side=LEFT, padx=5)

        if self.role == "operator":
            ctk.CTkButton(info_frame, text="🔒 SUBMIT & LOCK", width=130, height=30, fg_color="#D84315", hover_color="#BF360C", font=("Segoe UI", 11, "bold"), command=self.submit_and_lock).pack(side=LEFT, padx=20)

        self.double_label = ctk.CTkLabel(info_frame, text="DOUBLE : 0", font=("Segoe UI", 14, "bold"))
        self.double_label.pack(side=RIGHT, padx=10)
        self.normal_label = ctk.CTkLabel(info_frame, text="NORMAL : 0", font=("Segoe UI", 14, "bold"))
        self.normal_label.pack(side=RIGHT, padx=10)
        self.total_label = ctk.CTkLabel(info_frame, text="TOTAL : 0", font=("Segoe UI", 14, "bold"))
        self.total_label.pack(side=RIGHT, padx=10)

        style = ttk.Style()
        style.theme_use("default")
        style.configure("Treeview", rowheight=30, font=("Segoe UI", 11))
        style.configure("Treeview.Heading", font=("Segoe UI", 12, "bold"))
        
        columns = ("NO", "UNIQUE ID", "BARCODE", "STATUS", "USERNAME", "KETERANGAN DOUBLE", "WAKTU")
        self.tree = ttk.Treeview(self.root, columns=columns, show="headings")
        
        scroll_y = ctk.CTkScrollbar(self.root, orientation="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=scroll_y.set)
        scroll_y.pack(side=RIGHT, fill=Y, pady=10)

        for col in columns:
            self.tree.heading(col, text=col)
            self.tree.column(col, anchor="center")

        self.tree.column("NO", width=50)
        self.tree.column("UNIQUE ID", width=180)
        self.tree.column("BARCODE", width=200)

        self.tree.pack(fill=BOTH, expand=True, padx=10, pady=10)
        self.tree.tag_configure("double", background="#FFCDD2", foreground="#B71C1C")

        bottom_frame = ctk.CTkFrame(self.root, fg_color="#212121", height=35, corner_radius=0)
        bottom_frame.pack(fill=X, side=BOTTOM)

        self.status_label = ctk.CTkLabel(bottom_frame, text=" READY", font=("Segoe UI", 12, "bold"), text_color="white", anchor="w")
        self.status_label.pack(fill=Y, side=LEFT, padx=10)

        self.keep_focus()
        self.process_queue()
        
        self.load_table(date_filter=self.current_date_filter)

        if self.role in ["admin", "operator"]:
            self.root.bind("<F1>", lambda e: self.scan_entry.focus_force())
        if self.role in ["admin", "user"]:
            self.root.bind("<F2>", lambda e: self.search_data())
            self.root.bind("<F3>", lambda e: self.export_excel())

    # Fitur Baru: Window GUI untuk Manajemen Reminder / Catatan antar-user
    def open_reminders(self):
        win = ctk.CTkToplevel(self.root)
        win.title("REMINDER BERSAMA (CATATAN)")
        win.geometry("550x450")
        self.popup_config(win)

        ctk.CTkLabel(win, text="RIWAYAT REMINDER / CATATAN SYSTEM", font=("Segoe UI", 14, "bold")).pack(pady=10)

        tree_frame = ctk.CTkFrame(win)
        tree_frame.pack(fill=BOTH, expand=True, padx=15, pady=5)

        columns = ("USER", "PESAN CATATAN", "WAKTU")
        rem_tree = ttk.Treeview(tree_frame, columns=columns, show="headings")
        rem_tree.heading("USER", text="Oleh User")
        rem_tree.heading("PESAN CATATAN", text="Isi Catatan / Pengingat")
        rem_tree.heading("WAKTU", text="Waktu Dibuat")

        rem_tree.column("USER", width=90, anchor="center")
        rem_tree.column("PESAN CATATAN", width=300, anchor="w")
        rem_tree.column("WAKTU", width=120, anchor="center")

        scroll_y = ctk.CTkScrollbar(tree_frame, orientation="vertical", command=rem_tree.yview)
        rem_tree.configure(yscrollcommand=scroll_y.set)
        scroll_y.pack(side=RIGHT, fill=Y)
        rem_tree.pack(fill=BOTH, expand=True)

        def load_reminders_data():
            rem_tree.delete(*rem_tree.get_children())
            try:
                self.user_cursor.execute("SELECT username, pesan, waktu FROM reminders ORDER BY id DESC")
                for row in self.user_cursor.fetchall():
                    rem_tree.insert("", END, values=row)
            except Exception as e:
                print(f"Gagal memuat catatan: {e}")

        load_reminders_data()

        input_frame = ctk.CTkFrame(win, fg_color="transparent")
        input_frame.pack(fill=X, padx=15, pady=15)

        note_ent = ctk.CTkEntry(input_frame, placeholder_text="Tulis pengingat baru di sini...", font=("Segoe UI", 12), width=380)
        note_ent.pack(side=LEFT, padx=(0, 5), fill=X, expand=True)

        def save_new_reminder():
            msg = note_ent.get().strip()
            if not msg:
                messagebox.showwarning("WARNING", "Isi catatan tidak boleh kosong!")
                return
            try:
                now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                self.user_cursor.execute("INSERT INTO reminders(username, pesan, waktu) VALUES(?,?,?)", (self.current_user, msg, now_str))
                self.user_conn.commit()
                note_ent.delete(0, END)
                load_reminders_data()
            except Exception as e:
                messagebox.showerror("ERROR", f"Gagal menyimpan catatan: {e}")

        ctk.CTkButton(input_frame, text="KIRIM NOTE", width=100, font=("Segoe UI", 11, "bold"), fg_color="#00796B", hover_color="#004D40", command=save_new_reminder).pack(side=RIGHT)

    def submit_and_lock(self):
        target_date = self.current_date_filter if self.current_date_filter else datetime.now().strftime("%Y-%m-%d")
        confirm = messagebox.askyesno(
            "KONFIRMASI SUBMIT", 
            f"Apakah Anda yakin ingin Submit & Mengunci data pada tanggal {target_date}?\n\nSetelah dikunci, Operator tidak bisa menambah data tanpa izin Admin."
        )
        if not confirm: return
        
        try:
            self.cursor.execute("INSERT OR IGNORE INTO locked_dates(date) VALUES(?)", (target_date,))
            self.conn.commit()
            self.add_log(f"SUBMIT & LOCK Tanggal {target_date}")
            messagebox.showinfo("SUKSES", f"Data tanggal {target_date} berhasil di-submit dan dikunci.\nSistem akan otomatis keluar menuju Halaman Login.")
            
            self.logout()
        except Exception as e:
            messagebox.showerror("ERROR", str(e))

    def open_calendar_filter(self):
        cal_win = ctk.CTkToplevel(self.root)
        cal_win.title("PILIH TANGGAL")
        cal_win.geometry("350x400") 
        self.popup_config(cal_win)

        try:
            self.cursor.execute("SELECT DISTINCT SUBSTR(waktu, 1, 10) FROM scans")
            dates_with_data = [row[0] for row in self.cursor.fetchall() if row[0]]
        except Exception:
            dates_with_data = []

        # BUG FIX: Mengubah pattern menjadi 'yyyy-mm-dd' agar sesuai dengan standar string SQLite
        cal = Calendar(cal_win, selectmode='day', date_pattern='yyyy-mm-dd', showweeknumbers=False)
        cal.pack(pady=15, padx=20, fill=BOTH, expand=True)

        for date_str in dates_with_data:
            try:
                date_obj = datetime.strptime(date_str, "%Y-%m-%d").date()
                cal.calevent_create(date_obj, 'Data Ada', 'has_data')
            except Exception:
                pass

        cal.tag_config('has_data', background='#81C784', foreground='black')

        btn_frame = ctk.CTkFrame(cal_win, fg_color="transparent")
        btn_frame.pack(pady=10)

        def set_today():
            cal.selection_set(datetime.now().date())

        def apply_filter():
            selected_date = cal.get_date()
            self.current_date_filter = selected_date
            self.date_lbl.configure(text=self.current_date_filter)
            self.load_table(date_filter=self.current_date_filter)
            cal_win.destroy()

        ctk.CTkButton(btn_frame, text="📅 HARI INI", width=100, fg_color="#0288D1", hover_color="#01579B", font=("Segoe UI", 11, "bold"), command=set_today).pack(side=LEFT, padx=5)
        ctk.CTkButton(btn_frame, text="TERAPKAN", width=100, font=("Segoe UI", 11, "bold"), command=apply_filter).pack(side=LEFT, padx=5)
        
    def reset_date_filter(self):
        self.current_date_filter = ""
        self.date_lbl.configure(text="Semua Data")
        self.load_table()

    def process_queue(self):
        try:
            if not self.processing_scan and not self.scan_queue.empty():
                self.processing_scan = True
                barcode = self.scan_queue.get()
                self.process_barcode(barcode)
                self.processing_scan = False
        except Exception as e:
            print(f"Queue processing error: {e}")
        self.root.after(50, self.process_queue)

    def keep_focus(self):
        try:
            # BUG FIX: Menambahkan pengecekan 'self.scan_entry is not None' agar tidak memicu phantom loop saat dilogout
            if (self.is_running and not self.popup_open and hasattr(self, "scan_entry") 
                and self.scan_entry is not None and self.scan_entry.winfo_exists() and self.root.focus_get() != self.scan_entry):
                self.scan_entry.focus_force()
        except Exception:
            pass
        if self.is_running and hasattr(self, "scan_entry") and self.scan_entry is not None:
            self.scan_job = self.root.after(500, self.keep_focus)

    def request_admin_permission(self, target_date):
        approved = [False]
        win = ctk.CTkToplevel(self.root)
        win.title("OTORISASI IZIN ADMIN")
        win.geometry("380x290")
        self.popup_config(win)

        ctk.CTkLabel(win, text=f"⚠ PERINGATAN AKSES ⚠\nAnda mencoba menambah data pada tanggal ({target_date})\nyang status datanya telah DI-LOCK / DI-SUBMIT.", 
                     font=("Segoe UI", 11, "bold"), text_color="#B71C1C").pack(pady=10)
        
        ctk.CTkLabel(win, text="Username Admin:").pack(pady=2)
        user_entry = ctk.CTkEntry(win, width=240)
        user_entry.pack(pady=5)
        user_entry.focus_set()

        ctk.CTkLabel(win, text="Password Admin:").pack(pady=2)
        pass_entry = ctk.CTkEntry(win, show="*", width=240)
        pass_entry.pack(pady=5)

        def verify():
            usr = user_entry.get().strip().lower()
            pw = pass_entry.get()
            hashed_pw = hash_password(pw)
            
            self.user_cursor.execute("SELECT id FROM users WHERE username=? AND role='admin' AND password=?", (usr, hashed_pw))
            if self.user_cursor.fetchone():
                approved[0] = True
                win.destroy()
            else:
                messagebox.showerror("OTORISASI GAGAL", "Username / Password Salah atau User Tidak Memiliki Hak Admin!")

        ctk.CTkButton(win, text="BERI IZIN (VERIFIKASI)", fg_color="#388E3C", hover_color="#2E7D32", command=verify).pack(pady=15)
        pass_entry.bind("<Return>", lambda e: verify())
        
        self.root.wait_window(win)
        return approved[0]

    def scan_barcode(self):
        try:
            if not self.conn: return
            now = time.time()
            if now - self.last_scan_time < 0.2: return
            
            self.last_scan_time = now
            barcode = self.scan_entry.get().strip().replace("\n", "").replace("\r", "")
            
            if not barcode: return
            if len(barcode) < 11:
                messagebox.showerror("BARCODE ERROR", "Barcode minimal 11 digit")
                self.scan_entry.delete(0, END)
                return

            today_str = self.current_date_filter if self.current_date_filter else datetime.now().strftime("%Y-%m-%d")
            
            if self.role == "operator" and today_str not in self.session_authorized_dates:
                self.cursor.execute("SELECT 1 FROM locked_dates WHERE date=?", (today_str,))
                is_locked = self.cursor.fetchone() is not None
                
                if is_locked:
                    if not self.request_admin_permission(today_str):
                        self.scan_entry.delete(0, END)
                        return
                
                self.session_authorized_dates.add(today_str)

            self.scan_queue.put(barcode)
            self.add_log(f"SCAN {barcode}")
            self.scan_entry.delete(0, END)
        except Exception as e:
            print(f"Error in scan_barcode: {e}")

    def process_barcode(self, barcode):
        try:
            if self.current_date_filter:
                waktu = f"{self.current_date_filter} {datetime.now().strftime('%H:%M:%S')}"
                today_str = self.current_date_filter
            else:
                waktu = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                today_str = waktu[:10]

            self.cursor.execute(
                "SELECT id, unique_id, username, waktu FROM scans WHERE barcode=? AND waktu LIKE ? ORDER BY id ASC LIMIT 1", 
                (barcode, f"{today_str}%")
            )
            first_data = self.cursor.fetchone()

            status = "NORMAL"
            ket = ""
            tags = ()

            if first_data is not None:
                status = "DOUBLE"
                ref_uid = first_data[1] if first_data[1] else f"LEGACY-{first_data[0]}"
                ket = f"⚠ DOUBLE DENGAN ID {ref_uid} | USER {first_data[2]} | {first_data[3]}"
                tags = ("double",)

            db_prefix = os.path.basename(self.current_db).replace('.db', '').upper()
            db_code = ''.join(filter(str.isalnum, db_prefix))[:3].ljust(3, 'X') 
            date_compact = today_str.replace("-", "")
            
            self.cursor.execute("SELECT COUNT(*) FROM scans WHERE waktu LIKE ?", (f"{today_str}%",))
            seq = self.cursor.fetchone()[0] + 1
            unique_id = f"{db_code}-{date_compact}-{seq:04d}"

            self.cursor.execute(
                "INSERT INTO scans(barcode, status, username, waktu, unique_id) VALUES(?,?,?,?,?)",
                (barcode, status, self.current_user, waktu, unique_id)
            )
            self.conn.commit()
            row_id = self.cursor.lastrowid

            if not self.current_date_filter or waktu.startswith(self.current_date_filter):
                idx = self.total_count + 1
                self.tree.insert("", 0, iid=row_id, values=(idx, unique_id, barcode, status, self.current_user, ket, waktu), tags=tags)
                
                self.total_count += 1
                if status == "DOUBLE":
                    self.double_count += 1
                else:
                    self.normal_count += 1
                self.update_counter_labels()
            
            if status == "DOUBLE":
                self.status_label.configure(text=f" ⚠ DETECTED DOUBLE : {barcode} (Tetap Tersimpan)", text_color="#FFEB3B")
            else:
                self.status_label.configure(text=f" SCAN SUCCESS : {barcode}", text_color="white")
            
            if hasattr(self, "scan_entry") and self.scan_entry is not None:
                self.scan_entry.focus_force()

        except Exception as e:
            with open("error.log", "a", encoding="utf-8") as f:
                f.write(traceback.format_exc() + "\n")
            messagebox.showerror("SCAN ERROR", str(e))

    def update_counter_labels(self):
        try:
            self.total_label.configure(text=f"TOTAL : {self.total_count}")
            self.normal_label.configure(text=f"NORMAL : {self.normal_count}")
            self.double_label.configure(text=f"DOUBLE : {self.double_count}")
        except Exception:
            pass

    def load_table(self, date_filter=None):
        try:
            self.tree.delete(*self.tree.get_children())
            gc.collect() 
            
            if date_filter:
                self.cursor.execute("SELECT id, unique_id, barcode, status, username, waktu FROM scans WHERE waktu LIKE ? ORDER BY id ASC", (f"{date_filter}%",))
            else:
                self.cursor.execute("SELECT id, unique_id, barcode, status, username, waktu FROM scans ORDER BY id ASC")
                
            data = self.cursor.fetchall()

            self.total_count = 0
            self.normal_count = 0
            self.double_count = 0

            first_occurrences = {}

            for idx, row in enumerate(data, start=1):
                row_id, unique_id, barcode, db_status, username, waktu = row
                if not unique_id: unique_id = f"LEGACY-{row_id}"

                date_str = waktu[:10]
                key = f"{barcode}_{date_str}"

                ket = ""
                tags = ()

                if key in first_occurrences:
                    status = "DOUBLE"
                    fb = first_occurrences[key]
                    ket = f"⚠ DOUBLE DENGAN ID {fb['unique_id']} | {fb['user']} | {fb['time']}"
                    tags = ("double",)
                    self.double_count += 1
                else:
                    status = "NORMAL"
                    first_occurrences[key] = {"unique_id": unique_id, "user": username, "time": waktu}
                    self.normal_count += 1

                self.tree.insert("", 0, iid=row_id, values=(idx, unique_id, barcode, status, username, ket, waktu), tags=tags)
                self.total_count += 1

            self.update_counter_labels()
        except Exception as e:
            messagebox.showerror("LOAD ERROR", str(e))

    def export_excel(self):
        try:
            file = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")])
            if not file: return

            self.status_label.configure(text=" SEDANG MEMBUAT EXCEL... MOHON TUNGGU", text_color="yellow")

            if self.current_date_filter:
                self.cursor.execute("SELECT id, unique_id, barcode, status, username, waktu FROM scans WHERE waktu LIKE ? ORDER BY id ASC", (f"{self.current_date_filter}%",))
            else:
                self.cursor.execute("SELECT id, unique_id, barcode, status, username, waktu FROM scans ORDER BY id ASC")
            data = self.cursor.fetchall()

            if not data:
                messagebox.showwarning("WARNING", "Data kosong")
                self.status_label.configure(text=" READY", text_color="white")
                return

            db_name = os.path.basename(self.current_db).replace('.db', '').upper()
            threading.Thread(target=self._process_export_excel, args=(file, data, db_name), daemon=True).start()
        except Exception as e:
            messagebox.showerror("EXPORT ERROR", str(e))
            self.status_label.configure(text=" READY", text_color="white")

    def _process_export_excel(self, file, data, db_name):
        try:
            export_rows = []
            first_occurrences = {}

            for idx, row in enumerate(data, start=1):
                row_id, unique_id, barcode, db_status, username, waktu = row
                if not unique_id: unique_id = f"LEGACY-{row_id}"
                
                date_str = waktu[:10]
                key = f"{barcode}_{date_str}"
                
                ket = ""
                status = "NORMAL"
                
                if key in first_occurrences:
                    status = "DOUBLE"
                    fb = first_occurrences[key]
                    ket = f"⚠ DOUBLE DENGAN ID {fb['unique_id']} | {fb['user']} | {fb['time']}"
                else:
                    first_occurrences[key] = {"unique_id": unique_id, "user": username, "time": waktu}
                
                export_rows.append([idx, unique_id, barcode, status, ket, username, db_name, waktu])

            df = pd.DataFrame(export_rows, columns=["NO", "UNIQUE ID", "BARCODE", "STATUS", "KETERANGAN DOUBLE", "USER", "NAMA DATABASE", "WAKTU"])
            df.to_excel(file, index=False)
            
            try:
                import openpyxl
                from openpyxl.styles import PatternFill, Font
                wb = openpyxl.load_workbook(file)
                ws = wb.active
                red_fill = PatternFill(start_color="FFCDD2", end_color="FFCDD2", fill_type="solid")
                red_font = Font(color="B71C1C", bold=True)
                
                for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=8):
                    if row[3].value == "DOUBLE":
                        for cell in row:
                            cell.fill = red_fill
                            cell.font = red_font
                wb.save(file)
            except Exception as e:
                print(f"Error styling Excel: {e}")

            self.root.after(0, lambda: messagebox.showinfo("SUKSES", f"Export Excel Sukses dengan Pewarnaan otomatis\n{file}"))
            self.root.after(0, lambda: self.status_label.configure(text=" EXCEL BERHASIL DIBUAT", text_color="green"))
        except Exception as e:
            self.root.after(0, lambda: messagebox.showerror("EXPORT ERROR", str(e)))
            self.root.after(0, lambda: self.status_label.configure(text=" READY", text_color="white"))

    def export_pdf(self):
        try:
            file = filedialog.asksaveasfilename(
                defaultextension=".pdf", 
                filetypes=[("PDF Files", "*.pdf")],
                initialfile=f"EXPORT_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
            )
            if not file: return
            
            self.status_label.configure(text=" SEDANG MEMBUAT PDF... MOHON TUNGGU", text_color="yellow")

            if self.current_date_filter:
                self.cursor.execute("SELECT id, unique_id, barcode, status, username, waktu FROM scans WHERE waktu LIKE ? ORDER BY id ASC", (f"{self.current_date_filter}%",))
            else:
                self.cursor.execute("SELECT id, unique_id, barcode, status, username, waktu FROM scans ORDER BY id ASC")
            data = self.cursor.fetchall()
            
            if not data:
                messagebox.showwarning("WARNING", "Data kosong")
                self.status_label.configure(text=" READY", text_color="white")
                return

            db_name = os.path.basename(self.current_db).replace('.db', '').upper()
            threading.Thread(target=self._process_export_pdf, args=(file, data, db_name), daemon=True).start()
        except Exception as e:
            messagebox.showerror("PDF ERROR", str(e))
            self.status_label.configure(text=" READY", text_color="white")

    def _process_export_pdf(self, file, data, db_name):
        try:
            pdf = canvas.Canvas(file, pagesize=A4)
            width, height = A4
            y = height - 50 
            
            pdf.setFont("Helvetica-Bold", 14)
            pdf.drawString(50, y, f"LAPORAN SCAN RESI - DATABASE: {db_name}")
            y -= 30
            
            pdf.setFont("Helvetica", 7)
            
            first_occurrences = {}

            for idx, row in enumerate(data, start=1):
                row_id, unique_id, barcode, db_status, username, waktu = row
                if not unique_id: unique_id = f"LEGACY-{row_id}"
                
                date_str = waktu[:10]
                key = f"{barcode}_{date_str}"
                
                ket = ""
                status = "NORMAL"
                
                if key in first_occurrences:
                    status = "DOUBLE"
                    fb_uid = first_occurrences[key]
                    ket = f"DOUBLE ID {fb_uid}"
                    pdf.setFillColorRGB(0.8, 0, 0)
                else:
                    first_occurrences[key] = unique_id
                    pdf.setFillColorRGB(0, 0, 0)
                    
                line_text = f"No: {idx} | ID: {unique_id} | Resi: {barcode} | Status: {status} | Ket: {ket} | User: {username} | {waktu}"
                pdf.drawString(50, y, line_text)
                y -= 15
                if y < 50:
                    pdf.showPage()
                    y = height - 50
                    pdf.setFont("Helvetica", 7)
                    
            pdf.save()
            self.root.after(0, lambda: messagebox.showinfo("PDF SUKSES", f"PDF berhasil disimpan di\n{file}"))
            self.root.after(0, lambda: self.status_label.configure(text=" PDF BERHASIL DIBUAT", text_color="green"))
        except Exception as e:
            self.root.after(0, lambda: messagebox.showerror("PDF ERROR", str(e)))
            self.root.after(0, lambda: self.status_label.configure(text=" READY", text_color="white"))

    def import_excel(self):
        try:
            file = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx")])
            if not file: return
            
            df = pd.read_excel(file)
            for index, row in df.iterrows():
                barcode = str(row.iloc[0]).strip()
                waktu = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                today_str = waktu[:10]
                
                self.cursor.execute("SELECT id FROM scans WHERE barcode=? AND waktu LIKE ?", (barcode, f"{today_str}%"))
                status = "DOUBLE" if self.cursor.fetchone() else "NORMAL"
                
                unique_id = f"IMPORT-{datetime.now().strftime('%Y%m%d')}-{index+1:04d}"

                self.cursor.execute("INSERT INTO scans(barcode,status,username,waktu,unique_id) VALUES(?,?,?,?,?)",
                    (barcode, status, self.current_user, waktu, unique_id))
            self.conn.commit()
            self.load_table(date_filter=self.current_date_filter)
            messagebox.showinfo("IMPORT", "Import berhasil")
        except Exception as e:
            messagebox.showerror("IMPORT ERROR", str(e))

    def delete_data(self):
        try:
            selected = self.tree.selection()
            if not selected: return
            row_id = selected[0]
            self.cursor.execute("DELETE FROM scans WHERE id=?", (row_id,))
            self.conn.commit()
            self.load_table(date_filter=self.current_date_filter)
        except Exception as e:
            messagebox.showerror("DELETE ERROR", str(e))

    def edit_barcode(self):
        selected = self.tree.selection()
        if not selected: return

        item = self.tree.item(selected[0])
        row_id = selected[0]
        old_barcode = item["values"][2]

        win = ctk.CTkToplevel(self.root)
        win.title("EDIT BARCODE")
        win.geometry("350x180")
        self.popup_config(win)

        ctk.CTkLabel(win, text="BARCODE BARU").pack(pady=10)
        ent = ctk.CTkEntry(win, font=("Segoe UI", 14), width=250)
        ent.pack(padx=10)
        ent.insert(0, old_barcode)

        def save_edit():
            try:
                new_barcode = ent.get().strip()
                if len(new_barcode) < 11:
                    messagebox.showerror("ERROR", "Barcode minimal 11 digit")
                    return
                self.cursor.execute("UPDATE scans SET barcode=? WHERE id=?", (new_barcode, row_id))
                self.conn.commit()
                self.load_table(date_filter=self.current_date_filter)
                win.destroy()
            except Exception as e:
                messagebox.showerror("EDIT ERROR", str(e))

        ctk.CTkButton(win, text="SAVE", command=save_edit).pack(pady=20)

    def add_user(self):
        win = ctk.CTkToplevel(self.root)
        win.title("ADD USER")
        win.geometry("320x450")
        self.popup_config(win)

        ctk.CTkLabel(win, text="USERNAME BARU", font=("Segoe UI", 11, "bold")).pack(pady=(15, 0))
        user_entry = ctk.CTkEntry(win, width=250)
        user_entry.pack()

        ctk.CTkLabel(win, text="PASSWORD USER BARU", font=("Segoe UI", 11, "bold")).pack(pady=(10, 0))
        pw_entry = ctk.CTkEntry(win, show="*", width=250)
        pw_entry.pack()

        ctk.CTkLabel(win, text="ROLE / HAK AKSES", font=("Segoe UI", 11, "bold")).pack(pady=(10, 0))
        role_entry = ctk.CTkOptionMenu(win, values=["user", "operator", "admin"], width=250)
        role_entry.pack(pady=5)

        ctk.CTkLabel(win, text="--- OTORISASI ---", font=("Segoe UI", 10)).pack(pady=(15, 0))
        
        ctk.CTkLabel(win, text="PASSWORD ADMIN (ANDA)", font=("Segoe UI", 11, "bold")).pack(pady=(5, 0))
        admin_verif_entry = ctk.CTkEntry(win, show="*", width=250)
        admin_verif_entry.pack()

        def save():
            new_username = user_entry.get().strip().lower() 
            new_password = pw_entry.get()
            selected_role = role_entry.get()
            admin_pw_input = admin_verif_entry.get()

            if not new_username or not new_password or not admin_pw_input:
                messagebox.showwarning("PERINGATAN", "Semua kolom wajib diisi!")
                return

            try:
                hashed_admin_input = hash_password(admin_pw_input)
                self.user_cursor.execute("SELECT id FROM users WHERE username=? AND password=?", (self.current_user, hashed_admin_input))
                if not self.user_cursor.fetchone():
                    messagebox.showerror("ERROR", "Password Admin (Anda) Salah! Otorisasi Ditolak.")
                    return

                self.user_cursor.execute("INSERT INTO users(username,password,role) VALUES(?,?,?)",
                    (new_username, hash_password(new_password), selected_role))
                self.user_conn.commit()
                
                self.add_log(f"Membuat akun baru: {new_username} ({selected_role})")
                
                messagebox.showinfo("SUKSES", f"Akun '{new_username}' berhasil dibuat sebagai {selected_role.upper()}.")
                win.destroy()
            except sqlite3.IntegrityError:
                messagebox.showerror("ERROR", "Username tersebut sudah terdaftar!")
            except Exception as e:
                messagebox.showerror("ADD USER ERROR", str(e))

        ctk.CTkButton(win, text="BUAT AKUN", fg_color="#388E3C", hover_color="#2E7D32", command=save).pack(pady=20)

    def change_password(self):
        win = ctk.CTkToplevel(self.root)
        win.title("GANTI PASSWORD USER")
        win.geometry("350x380")
        self.popup_config(win)

        ctk.CTkLabel(win, text="PILIH USER", font=("Segoe UI", 12, "bold")).pack(pady=(15, 5))

        try:
            self.user_cursor.execute("SELECT username FROM users")
            users = [row[0] for row in self.user_cursor.fetchall()]
        except Exception as e:
            messagebox.showerror("ERROR", str(e))
            win.destroy()
            return

        selected_user = StringVar(value=users[0])
        dropdown = ctk.CTkOptionMenu(win, variable=selected_user, values=users, width=250)
        dropdown.pack(pady=5)

        ctk.CTkLabel(win, text="PASSWORD BARU", font=("Segoe UI", 12, "bold")).pack(pady=(10, 5))
        new_pw = ctk.CTkEntry(win, show="*", width=250)
        new_pw.pack()

        ctk.CTkLabel(win, text="KONFIRMASI PASSWORD", font=("Segoe UI", 12, "bold")).pack(pady=(10, 5))
        confirm_pw = ctk.CTkEntry(win, show="*", width=250)
        confirm_pw.pack()

        def save_new_password():
            npw = new_pw.get()
            cpw = confirm_pw.get()
            usr = selected_user.get()

            if not npw or not cpw:
                messagebox.showerror("ERROR", "Password tidak boleh kosong!")
                return
            if npw != cpw:
                messagebox.showerror("ERROR", "Password dan Konfirmasi tidak cocok!")
                return

            try:
                hashed = hash_password(npw)
                self.user_cursor.execute("UPDATE users SET password=? WHERE username=?", (hashed, usr))
                self.user_conn.commit()
                
                self.add_log(f"Ubah password user: {usr}")
                
                messagebox.showinfo("SUKSES", f"Password untuk user '{usr}' berhasil diubah!")
                win.destroy()
            except Exception as e:
                messagebox.showerror("ERROR", f"Gagal mengubah password:\n{str(e)}")

        ctk.CTkButton(win, text="SIMPAN PASSWORD", fg_color="#388E3C", hover_color="#2E7D32", command=save_new_password).pack(pady=25)

    def search_data(self):
        win = ctk.CTkToplevel(self.root)
        win.title("SEARCH")
        win.geometry("350x180")
        self.popup_config(win)

        ctk.CTkLabel(win, text="BARCODE").pack(pady=5)
        ent = ctk.CTkEntry(win, font=("Segoe UI", 14), width=250)
        ent.pack(padx=10)

        def do_search():
            try:
                key = ent.get().strip()
                self.tree.delete(*self.tree.get_children())
                
                if self.current_date_filter:
                    self.cursor.execute("SELECT id, unique_id, barcode, status, username, waktu FROM scans WHERE barcode LIKE ? AND waktu LIKE ? ORDER BY id ASC", (f"%{key}%", f"{self.current_date_filter}%"))
                else:
                    self.cursor.execute("SELECT id, unique_id, barcode, status, username, waktu FROM scans WHERE barcode LIKE ? ORDER BY id ASC", (f"%{key}%",))
                    
                results = self.cursor.fetchall()
                
                # BUG FIX: Dioptimalkan menggunakan hashmap (dictionary) di memori agar pencarian berkecepatan tinggi ($O(1)$)
                # dan mencegah penimpangan cursor state di dalam perulangan loop.
                first_occurrences = {}
                for idx, row in enumerate(results, 1):
                    row_id, unique_id, barcode, db_status, username, waktu = row
                    if not unique_id: unique_id = f"LEGACY-{row_id}"
                    
                    date_str = waktu[:10]
                    map_key = f"{barcode}_{date_str}"
                    
                    ket, tags, status = ("", (), "NORMAL")
                    if map_key in first_occurrences:
                        ref_fb = first_occurrences[map_key]
                        ket = f"⚠ DOUBLE DENGAN ID {ref_fb['unique_id']} | {ref_fb['user']} | {ref_fb['time']}"
                        status = "DOUBLE"
                        tags = ("double",)
                    else:
                        first_occurrences[map_key] = {"unique_id": unique_id, "user": username, "time": waktu}
                    
                    self.tree.insert("", END, iid=row_id, values=(idx, unique_id, barcode, status, username, ket, waktu), tags=tags)
            except Exception as e:
                messagebox.showerror("SEARCH ERROR", str(e))

        ctk.CTkButton(win, text="SEARCH", command=do_search).pack(pady=(10, 5))
        ctk.CTkButton(win, text="RESET TABLE", fg_color="gray", command=lambda: [self.load_table(date_filter=self.current_date_filter), win.destroy()]).pack()

    def view_logs(self):
        win = ctk.CTkToplevel(self.root)
        win.title("USER LOG")
        win.geometry("900x500")
        self.popup_config(win)

        tree = ttk.Treeview(win, columns=("USER", "ACTIVITY", "TIME"), show="headings")
        tree.heading("USER", text="USER")
        tree.heading("ACTIVITY", text="ACTIVITY")
        tree.heading("TIME", text="TIME")
        tree.pack(fill=BOTH, expand=True)

        self.cursor.execute("SELECT username,activity,waktu FROM logs ORDER BY id DESC")
        for row in self.cursor.fetchall():
            tree.insert("", END, values=row)

    def backup_database(self):
        try:
            if not self.conn: return
            
            db_basename = os.path.basename(self.current_db).replace('.db', '')
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            backup_filename = f"backup_{db_basename}_{timestamp}.db"
            
            save_path = os.path.join(BACKUP_FOLDER, backup_filename)

            self.conn.commit()
            with sqlite3.connect(save_path) as backup_conn:
                self.conn.backup(backup_conn)

            self.add_log(f"Backup Database: {backup_filename}")
            messagebox.showinfo("BACKUP BERHASIL", f"Database berhasil dicopy/dibackup otomatis di folder 'backup':\n\n{backup_filename}")
        except Exception as e:
            messagebox.showerror("BACKUP ERROR", str(e))

    def close_database(self):
        try:
            self.is_running = False
            if self.scan_job: self.root.after_cancel(self.scan_job)
            if self.conn:
                self.conn.commit()
                self.cursor.execute("PRAGMA wal_checkpoint(FULL)")
                self.conn.close()
        except Exception:
            pass
        self.conn = None
        self.cursor = None
        self.clear()
        self.database_selector_gui()

    # Fitur Baru: Otomatisasi validasi versi ke Server/GitHub API (Asynchronous Threading)
    def check_app_updates(self, silent=False):
        def async_check():
            try:
                # Masukkan URL JSON Config Update Asli Anda di sini nanti (Contoh menggunakan file JSON raw GitHub)
                update_url = "https://raw.githubusercontent.com/username/repo/main/version.json"
                
                req = urllib.request.Request(update_url, headers={'User-Agent': 'Mozilla/5.0'})
                with urllib.request.urlopen(req, timeout=5) as response:
                    data = json.loads(response.read().decode())
                    latest_version = data.get("version", VERSION)
                    download_link = data.get("url", "https://github.com")
                    
                    if latest_version > VERSION:
                        self.root.after(0, lambda: self.prompt_update_dialog(latest_version, download_link))
                    else:
                        if not silent:
                            self.root.after(0, lambda: messagebox.showinfo("UPDATE SYSTEM", f"Aplikasi Anda sudah versi terbaru!\nVersi Saat Ini: V{VERSION}"))
            except Exception as e:
                if not silent:
                    self.root.after(0, lambda: messagebox.showwarning("UPDATE GAGAL", f"Gagal memeriksa pembaruan.\nPeriksa koneksi internet Anda atau hubungi Developer.\n\nDetail: {e}"))
        
        threading.Thread(target=async_check, daemon=True).start()

    def prompt_update_dialog(self, latest_v, url):
        ans = messagebox.askyesno("UPDATE TERSEDIA", f"Versi Baru V{latest_v} Telah Tersedia!\n\nApakah Anda ingin mengunduh pembaruan/installer terbaru sekarang untuk troubleshooting?")
        if ans:
            webbrowser.open(url)

    def show_about_popup(self):
        about_win = ctk.CTkToplevel(self.root)
        about_win.title("ABOUT APPLICATION")
        about_win.geometry("460x520")
        about_win.resizable(False, False)
        self.popup_config(about_win)
        
        main_frame = ctk.CTkFrame(about_win, fg_color="#ECEFF1", corner_radius=15)
        main_frame.pack(fill=BOTH, expand=True, padx=15, pady=15)
        
        title_lbl = ctk.CTkLabel(main_frame, text=f"Resi Scanner V{VERSION}", font=("Segoe UI", 22, "bold"), text_color="#1A237E")
        title_lbl.pack(pady=(20, 15))
        
        def create_section(section_title, content_text):
            lbl_sec = ctk.CTkLabel(main_frame, text=section_title, font=("Segoe UI", 11, "bold"), text_color="#546E7A")
            lbl_sec.pack(pady=(6, 0))
            lbl_content = ctk.CTkLabel(main_frame, text=f'"{content_text}"', font=("Consolas", 11, "italic"), text_color="#263238", wraplength=400)
            lbl_content.pack(pady=(0, 6))
            
        create_section("Bahasa pemograman yang digunakan", "python")
        create_section("Jenis Engine yang di gunakan", "Sqllite,customtkinter, tkcalendar, pandas, openpyxl, reportlab")
        create_section("Troubleshooting", "please contact admin")
        create_section("thanks to", "Tuhan yang maha Esa, GitHub, W3schools, @ar1nben1, fromwire, & Papaya Fresh Gallery Bandung")
        
        footer_lbl = ctk.CTkLabel(main_frame, text="@2026\nseftian_permana. | hotline: (0857-5926-7590)", font=("Segoe UI", 11, "bold"), text_color="#37474F", justify="center")
        footer_lbl.pack(pady=(15, 10))
        
        btn_action_frame = ctk.CTkFrame(main_frame, fg_color="transparent")
        btn_action_frame.pack(pady=5)
        
        # Fitur Baru: Menambahkan tombol Cek Update di panel ABOUT
        update_btn = ctk.CTkButton(btn_action_frame, text="🔄 CEK UPDATE", width=120, height=30, fg_color="#E65100", hover_color="#BF360C", font=("Segoe UI", 11, "bold"), command=lambda: self.check_app_updates(silent=False))
        update_btn.pack(side=LEFT, padx=5)
        
        close_btn = ctk.CTkButton(btn_action_frame, text="TUTUP", width=100, height=30, fg_color="#5C6BC0", hover_color="#3949AB", font=("Segoe UI", 11, "bold"), command=about_win.destroy)
        close_btn.pack(side=LEFT, padx=5)

    def reset_session(self):
        try:
            self.is_running = False
            if self.scan_job: self.root.after_cancel(self.scan_job)
            if self.conn:
                self.conn.commit()
                self.conn.close()
        except Exception:
            pass
        self.conn = None
        self.cursor = None
        self.popup_open = False
        self.scan_job = None
        self.scan_entry = None  # BUG FIX: Reset nilai entry ke None untuk memutus looping focus ilegal
        self.current_date_filter = ""
        self.total_count = 0
        self.normal_count = 0
        self.double_count = 0
        self.session_authorized_dates = set()
        gc.collect()

    def logout(self):
        self.reset_session()
        self.current_user = ""
        self.role = ""
        self.clear()
        self.is_running = True
        self.login_ui()

    def exit_app(self):
        try:
            self.is_running = False
            if self.conn:
                self.conn.commit()
                self.cursor.execute("PRAGMA wal_checkpoint(FULL)")
                self.conn.close()
            if self.user_conn:
                self.user_conn.commit()
                self.user_conn.close()
        except Exception:
            pass
        try:
            self.root.quit()
            self.root.destroy()
        except Exception:
            os._exit(0)

if __name__ == "__main__":
    multiprocessing.freeze_support()
    root = ctk.CTk()
    root.resizable(True, True)
    root.minsize(1200, 700)
    app = App(root)
    root.mainloop()